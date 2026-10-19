from pathlib import Path
import csv
import openpyxl


def serialize_file(path: str) -> str:
    ext = Path(path).suffix.lower()

    if ext in ('.xlsx', '.xlsm'):
        return serialize_xlsx(path)
    elif ext == '.csv':
        return serialize_csv(csv)
    else:
        return ValueError(f"unsupported file type: {ext}")


def serialize_csv(path: str) -> str:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)

    lines = [f"# FILE: {Path(path).name}", "#SHEET: (csv, single table)", ""]

    for i, row in enumerate(rows):
        lines.append(f"R{i}: " + " | ".join(cell.strip() for cell in row))

        return "\n".join(lines)

def serialize_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    out = [f"# FILE: {Path(path).name}"]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        out.append(f"\n# SHEET: {sheet_name}")
        out.append(f"# DIMENSIONS: {ws.dimensions}")


        if ws.merged_cells.ranges:
            merged_str = ", ".join(str(r) for r in ws.merged_cells.ranges)
            out.append(f"# MERGED_CELLS: {merged_str}")
 
        out.append("")
 
        for row in ws.iter_rows():
            row_num = row[0].row
            cells = []
            for cell in row:
                val = cell.value
                if val is None:
                    cells.append("")
                else:
                    cells.append(str(val).strip())
            if all(c == "" for c in cells):
                out.append(f"R{row_num}: <BLANK ROW>")
            else:
                out.append(f"R{row_num}: " + " | ".join(cells))
 
    return "\n".join(out)

